"""Pruebas de la API REST: el mismo núcleo, ahora a través del adaptador JSON.

Cubren los códigos que pide la guía del Sprint: 200/201, 400, 401, 403, 404 y 409.
"""
from io import BytesIO

from openpyxl import load_workbook

from tests.conftest import token

HUGO, ANA, KARLA = 12, 5, 15  # ids de los datos de prueba


def _curso_de(cliente, cabeceras) -> int:
    return cliente.get("/api/cursos", headers=cabeceras).json()[0]["id"]


# ---------- HU-01: login y RBAC ----------

def test_login_entrega_token_jwt(cliente):
    r = cliente.post("/api/login", json={"usuario": "pjara", "clave": "profe123"})
    assert r.status_code == 200
    assert r.json()["token_type"] == "bearer" and r.json()["rol"] == "profesor"


def test_login_con_clave_incorrecta_es_401(cliente):
    r = cliente.post("/api/login", json={"usuario": "pjara", "clave": "mala"})
    assert r.status_code == 401
    assert r.json()["detail"] == "Usuario o contraseña incorrectos"


def test_sin_token_es_401(cliente):
    assert cliente.get("/api/cursos").status_code == 401


def test_token_falso_es_401(cliente):
    assert cliente.get("/api/cursos", headers={"Authorization": "Bearer abc.def.ghi"}).status_code == 401


def test_alumno_no_puede_ver_cursos_403(cliente):
    r = cliente.get("/api/cursos", headers=token(cliente, "alumno1", "alumno123"))
    assert r.status_code == 403


def test_profesor_no_ve_riesgo_de_curso_ajeno_403(cliente):
    cab = token(cliente, "cmunoz", "profe123")
    propio = _curso_de(cliente, cab)
    assert cliente.get(f"/api/cursos/{propio - 1}/riesgo", headers=cab).status_code == 403


def test_curso_inexistente_es_404(cliente):
    cab = token(cliente, "pjara", "profe123")
    assert cliente.get("/api/cursos/999/riesgo", headers=cab).status_code == 404


# ---------- HU-04 + HU-07 ----------

def test_riesgo_del_curso_con_semaforo(cliente):
    cab = token(cliente, "pjara", "profe123")
    r = cliente.get(f"/api/cursos/{_curso_de(cliente, cab)}/riesgo", headers=cab)
    assert r.status_code == 200
    datos = r.json()
    assert datos["alumnos"][0]["riesgo"] == "Alto" and datos["alumnos"][0]["semaforo"] == "Rojo"
    assert datos["conteo"]["Sin datos"] == 1


def test_la_directiva_no_ve_el_riesgo(cliente):
    cab = token(cliente, "directiva", "directiva123")
    assert cliente.get("/api/cursos/1/riesgo", headers=cab).status_code == 403


# ---------- HU-05 ----------

def test_guardar_notas_recalcula_el_riesgo(cliente):
    cab = token(cliente, "pjara", "profe123")
    curso = _curso_de(cliente, cab)
    r = cliente.put(f"/api/cursos/{curso}/notas", headers=cab, json={
        "asignatura_id": 1, "notas": [{"alumno_id": HUGO, "numero": 1, "nota": 3.0}]})
    assert r.status_code == 200
    hugo = next(a for a in cliente.get(f"/api/cursos/{curso}/riesgo", headers=cab).json()["alumnos"]
                if a["id"] == HUGO)
    assert hugo["riesgo"] == "Medio"


def test_nota_fuera_de_rango_es_400_y_no_guarda_nada(cliente):
    cab = token(cliente, "pjara", "profe123")
    curso = _curso_de(cliente, cab)
    r = cliente.put(f"/api/cursos/{curso}/notas", headers=cab, json={
        "asignatura_id": 1, "notas": [{"alumno_id": HUGO, "numero": 1, "nota": 6.0},
                                      {"alumno_id": HUGO, "numero": 2, "nota": 7.5}]})
    assert r.status_code == 400
    assert f"n_{HUGO}_2" in r.json()["errores"]
    assert cliente.get(f"/api/cursos/{curso}/notas?asignatura_id=1", headers=cab).json().get(str(HUGO)) is None


# ---------- HU-06 ----------

def test_pasar_lista_en_feriado_es_400(cliente):
    cab = token(cliente, "pjara", "profe123")
    r = cliente.put(f"/api/cursos/{_curso_de(cliente, cab)}/asistencia", headers=cab,
                    json={"fecha": "2026-09-18", "registros": {}})
    assert r.status_code == 400
    assert "Independencia Nacional" in r.json()["detail"]


def test_estado_de_asistencia_invalido_es_422(cliente):
    cab = token(cliente, "pjara", "profe123")
    r = cliente.put(f"/api/cursos/{_curso_de(cliente, cab)}/asistencia", headers=cab,
                    json={"fecha": "2026-10-06", "registros": {str(HUGO): "X"}})
    assert r.status_code == 422  # Pydantic valida el formato antes de llegar al núcleo


def test_pasar_lista_valida(cliente):
    cab = token(cliente, "pjara", "profe123")
    curso = _curso_de(cliente, cab)
    r = cliente.put(f"/api/cursos/{curso}/asistencia", headers=cab,
                    json={"fecha": "2026-10-06", "registros": {str(HUGO): "A"}})
    assert r.status_code == 200
    dia = cliente.get(f"/api/cursos/{curso}/asistencia?fecha=2026-10-06", headers=cab).json()
    assert dia["registros"][str(HUGO)] == "A"


# ---------- HU-02 ----------

def test_directiva_crea_curso_201_y_nombre_repetido_409(cliente):
    cab = token(cliente, "directiva", "directiva123")
    r = cliente.post("/api/cursos", headers=cab, json={"nombre": "4° Medio A", "profesor_id": 4})
    assert r.status_code == 201
    assert cliente.post("/api/cursos", headers=cab, json={"nombre": "4° Medio A"}).status_code == 409


def test_alumno_en_otro_curso_es_409(cliente):
    cab = token(cliente, "directiva", "directiva123")
    assert cliente.post("/api/cursos/2/alumnos", headers=cab, json={"alumno_id": ANA}).status_code == 409
    assert cliente.post("/api/cursos/2/alumnos", headers=cab, json={"alumno_id": KARLA}).status_code == 201


def test_profesor_no_puede_crear_cursos_403(cliente):
    cab = token(cliente, "pjara", "profe123")
    assert cliente.post("/api/cursos", headers=cab, json={"nombre": "X"}).status_code == 403


# ---------- HU-03 ----------

def test_semestre_invertido_es_400(cliente):
    cab = token(cliente, "directiva", "directiva123")
    r = cliente.put("/api/calendario/semestre", headers=cab, json={"inicio": "2026-12-01", "fin": "2026-08-01"})
    assert r.status_code == 400


def test_feriado_nuevo_y_eliminar(cliente):
    cab = token(cliente, "directiva", "directiva123")
    assert cliente.post("/api/calendario/feriados", headers=cab,
                        json={"fecha": "2026-10-06", "nombre": "Prueba"}).status_code == 201
    assert cliente.delete("/api/calendario/feriados/2026-10-06", headers=cab).status_code == 204
    assert cliente.delete("/api/calendario/feriados/2026-10-06", headers=cab).status_code == 404


# ---------- HU-08 ----------

def test_exportar_asistencia_xlsx(cliente):
    cab = token(cliente, "pjara", "profe123")
    r = cliente.get(f"/api/cursos/{_curso_de(cliente, cab)}/exportar/asistencia", headers=cab)
    assert r.status_code == 200
    ws = load_workbook(BytesIO(r.content)).active
    assert [c.value for c in ws[1]][-4:] == ["Presente", "Ausente", "Justificado", "% Asistencia"]


# ---------- Swagger ----------

def test_swagger_documenta_los_endpoints(cliente):
    rutas = cliente.get("/openapi.json").json()["paths"]
    for ruta in ("/api/login", "/api/cursos/{curso_id}/riesgo", "/api/cursos/{curso_id}/notas",
                 "/api/cursos/{curso_id}/asistencia", "/api/calendario/semestre"):
        assert ruta in rutas
