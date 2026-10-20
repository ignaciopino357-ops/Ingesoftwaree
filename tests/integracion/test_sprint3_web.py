"""Criterios de aceptación del Sprint 3 probados a través de las páginas web (base SQLite real)."""
from datetime import date, timedelta

from tests.conftest import entrar

# Ids de los datos de prueba (app/compartido/seed.py)
CURSO_A, CURSO_B = 1, 2
CAMILA, EMILIA, HUGO, ISIDORA = 7, 9, 12, 13


# ---------- HU-09 ----------

def test_alumno_ve_sus_notas_y_asistencia_sin_riesgo(cliente):
    entrar(cliente, "alumno3", "alumno123")   # Camila, 3° Medio A
    r = cliente.get("/alumno")
    assert r.status_code == 200
    assert "Camila" in r.text and "3° Medio A" in r.text
    assert "Lenguaje" in r.text and "asistencia" in r.text
    assert "Días que falté" in r.text and "02-10" in r.text
    assert "riesgo" not in r.text.lower() and "semáforo" not in r.text and 'class="punto' not in r.text


def test_alumno_no_ve_datos_de_otro(cliente):
    entrar(cliente, "alumno3", "alumno123")
    assert "Emilia" not in cliente.get("/alumno").text
    assert cliente.get(f"/profesor/cursos/{CURSO_A}/alumnos/{EMILIA}").status_code == 403


# ---------- HU-16 ----------

def test_calendario_del_alumno_pinta_ausencias_y_feriados(cliente):
    entrar(cliente, "alumno3", "alumno123")
    r = cliente.get("/alumno/asistencia?mes=2026-10")
    assert "Octubre 2026" in r.text
    assert 'class="ausente" title="Ausente">2<' in r.text           # 2 de octubre: ausente
    assert 'title="Es feriado: Encuentro de Dos Mundos">12<' in r.text
    assert "riesgo" not in r.text.lower()


def test_calendario_no_sale_del_semestre(cliente):
    entrar(cliente, "alumno3", "alumno123")
    r = cliente.get("/alumno/asistencia?mes=2026-07")
    assert "Julio 2026" in r.text and 'aria-label="Mes anterior"' not in r.text
    assert "Diciembre 2026" in cliente.get("/alumno/asistencia?mes=2030-01").text


# ---------- HU-10 y HU-12 ----------

def test_profesor_ve_alertas_iniciales_y_contador(cliente):
    entrar(cliente, "pjara", "profe123")
    r = cliente.get("/profesor/alertas")
    assert "Pasó a riesgo Alto" in r.text and "Emilia Vargas" in r.text
    assert "3 ausencias seguidas: 02-10, 05-10, 06-10" in r.text
    assert 'aria-label="2 pendientes"' in r.text


def test_marcar_alerta_revisada_la_saca_de_pendientes(cliente):
    entrar(cliente, "pjara", "profe123")
    r = cliente.post("/profesor/alertas/1/revisar", follow_redirects=True)
    assert "Alerta marcada como revisada" in r.text
    assert 'aria-label="1 pendientes"' in r.text


def test_otro_profesor_no_ve_ni_revisa_alertas_ajenas(cliente):
    entrar(cliente, "cmunoz", "profe123")
    assert "Emilia" not in cliente.get("/profesor/alertas").text
    assert cliente.post("/profesor/alertas/1/revisar").status_code == 403


def test_guardar_notas_que_dejan_en_alto_crea_alerta(cliente):
    entrar(cliente, "pjara", "profe123")
    # Camila está en Medio por asistencia; con notas rojas en todas las asignaturas pasa a Alto
    for asignatura in range(1, 6):
        datos = {"asignatura_id": str(asignatura)}
        datos.update({f"n_{CAMILA}_{n}": "2,0" for n in range(1, 4)})
        cliente.post(f"/profesor/cursos/{CURSO_A}/notas", data=datos)
    r = cliente.get("/profesor/alertas")
    assert r.text.count("Camila Fuentes") == 2           # ausencias + riesgo Alto
    assert r.text.count("Pasó a riesgo Alto") == 2       # Emilia y Camila
    # guardar otra vez no repite la alerta
    cliente.post(f"/profesor/cursos/{CURSO_A}/notas", data={"asignatura_id": "1", f"n_{CAMILA}_4": "2,5"})
    assert cliente.get("/profesor/alertas").text.count("Pasó a riesgo Alto") == 2


def test_pasar_lista_con_tercera_falta_crea_alerta(cliente):
    entrar(cliente, "cmunoz", "profe123")   # 3° Medio B: Isidora y Joaquín
    dias = []
    dia = date(2026, 10, 7)
    while len(dias) < 3:                     # 3 días de clases seguidos después de los datos de prueba
        if dia.weekday() < 5 and dia != date(2026, 10, 12):
            dias.append(dia)
        dia += timedelta(days=1)
    for d in dias:
        cliente.post(f"/profesor/cursos/{CURSO_B}/asistencia",
                     data={"fecha": d.isoformat(), f"estado_{ISIDORA}": "A"})
    r = cliente.get("/profesor/alertas")
    # Isidora ya había faltado el 06-10 en los datos de prueba: la racha parte ese día y
    # la alerta se crea una sola vez, aunque la racha siga creciendo el 09-10
    assert "Isidora Pérez" in r.text
    assert r.text.count("ausencias seguidas:") == 1 and "06-10, 07-10, 08-10" in r.text


# ---------- HU-11 ----------

def test_directiva_registra_alumno_con_curso_y_puede_entrar(cliente):
    entrar(cliente, "directiva", "directiva123")
    r = cliente.post("/directiva/usuarios", data={"nombre": "Nora Paz", "usuario": "npaz", "clave": "secreto1",
                                                  "rol": "alumno", "curso_id": str(CURSO_B)}, follow_redirects=True)
    assert "Alumno Nora Paz registrado" in r.text and "Quedó en 3° Medio B" in r.text
    assert "Nora Paz" in cliente.get("/directiva/usuarios").text
    cliente.post("/logout")
    assert entrar(cliente, "npaz", "secreto1").headers["location"] == "/alumno"


def test_registrar_profesor_aparece_en_lista_para_asignar_curso(cliente):
    entrar(cliente, "directiva", "directiva123")
    cliente.post("/directiva/usuarios", data={"nombre": "Tomás Ruiz", "usuario": "truiz", "clave": "secreto1",
                                              "rol": "profesor"})
    assert "Tomás Ruiz" in cliente.get(f"/directiva/cursos/{CURSO_A}").text


def test_usuario_repetido_no_se_guarda(cliente):
    entrar(cliente, "directiva", "directiva123")
    r = cliente.post("/directiva/usuarios", data={"nombre": "Otro", "usuario": "alumno1", "clave": "secreto1",
                                                  "rol": "alumno"})
    assert r.status_code == 409 and "el usuario alumno1 ya existe" in r.text
    assert 'value="Otro"' in r.text                       # el formulario conserva lo escrito


def test_clave_corta_marca_el_campo(cliente):
    entrar(cliente, "directiva", "directiva123")
    r = cliente.post("/directiva/usuarios", data={"nombre": "Nora", "usuario": "npaz", "clave": "123",
                                                  "rol": "alumno"})
    assert r.status_code == 400 and "al menos 6 caracteres" in r.text


def test_solo_directiva_registra_usuarios(cliente):
    entrar(cliente, "pjara", "profe123")
    assert cliente.get("/directiva/usuarios").status_code == 403
    assert cliente.post("/directiva/usuarios", data={"nombre": "X", "usuario": "xxx", "clave": "secreto1",
                                                     "rol": "alumno"}).status_code == 403


# ---------- HU-13 ----------

def test_historial_semanal_en_la_ficha_del_alumno(cliente):
    entrar(cliente, "pjara", "profe123")
    r = cliente.get(f"/profesor/cursos/{CURSO_A}/alumnos/{EMILIA}")
    assert "<title>Emilia Vargas · Seguimiento escolar</title>" in r.text
    assert "Historial de riesgo por semana" in r.text
    assert "<polyline" in r.text and "umbral 85 %" in r.text and "umbral 4.0" in r.text
    assert "27-07 al 02-08" in r.text                    # primera semana del semestre
    sin_datos = cliente.get(f"/profesor/cursos/{CURSO_A}/alumnos/{HUGO}")
    assert "Sin datos" in sin_datos.text.split("Historial de riesgo")[1]


# ---------- HU-14 ----------

def test_filtrar_curso_por_color(cliente):
    entrar(cliente, "pjara", "profe123")
    r = cliente.get(f"/profesor/cursos/{CURSO_A}?ver=Alto")
    tabla = r.text.split("<tbody>")[1]
    assert "Emilia Vargas" in tabla and "Ana Contreras" not in tabla
    assert "Mostrando 1 de 8 alumnos (Alto)" in r.text
    dos = cliente.get(f"/profesor/cursos/{CURSO_A}?ver=Alto&ver=Medio").text
    assert "Mostrando 3 de 8 alumnos" in dos


def test_filtro_no_cambia_la_exportacion(cliente):
    from io import BytesIO

    from openpyxl import load_workbook
    entrar(cliente, "pjara", "profe123")
    cliente.get(f"/profesor/cursos/{CURSO_A}?ver=Alto")
    hoja = load_workbook(BytesIO(cliente.get(f"/profesor/cursos/{CURSO_A}/exportar/notas").content)).active
    assert hoja.max_row - 1 >= 8                          # todos los alumnos del curso, no solo los filtrados


# ---------- HU-15 ----------

def test_resumen_de_riesgo_por_curso_sin_nombres(cliente):
    entrar(cliente, "directiva", "directiva123")
    r = cliente.get("/directiva/riesgo")
    assert "3° Medio A" in r.text and "Pedro Jara" in r.text
    for alumno in ("Emilia", "Camila", "Ana Contreras"):
        assert alumno not in r.text
    assert "10</strong> alumnos en cursos" in r.text


def test_profesor_no_ve_resumen_de_directiva(cliente):
    entrar(cliente, "pjara", "profe123")
    assert cliente.get("/directiva/riesgo").status_code == 403
