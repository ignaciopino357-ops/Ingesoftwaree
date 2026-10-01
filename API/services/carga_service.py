# API/services/carga_service.py
# Parser de Excel/CSV con validación de tipos de datos.
# Reglas: nota entre 1.0 y 7.0, asistencia entre 0 y 100, el alumno debe pertenecer al curso.
#
# Formato esperado (encabezados en cualquier orden, sin importar mayúsculas):
#   id_estudiante | fecha | nota | asistencia        (fecha: AAAA-MM-DD)
# Volver a subir el mismo alumno + fecha REEMPLAZA el registro anterior.

import csv
import io
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook

from API.database import conexion_bd
from API.excepciones import CursoNoEncontradoException

COLUMNAS_OBLIGATORIAS = {"id_estudiante", "fecha", "nota", "asistencia"}


class ArchivoInvalidoException(Exception):
    """Extensión no soportada o encabezados incorrectos: se rechaza antes de procesar."""


def _a_float(valor: Any) -> float:
    if isinstance(valor, str):
        valor = valor.strip().replace(",", ".")  # Excel en español usa coma decimal
    return float(valor)


class CargaService:

    # ---------- Lectura de archivo -> lista de (nro_fila, dict) ----------
    def leer_archivo(self, nombre_archivo: str, contenido: bytes) -> List[Tuple[int, Dict[str, Any]]]:
        nombre = (nombre_archivo or "").lower()
        extension = nombre.rsplit(".", 1)[-1] if "." in nombre else ""
        if extension == "csv":
            matriz = self._leer_csv(contenido)
        elif extension in ("xlsx", "xlsm"):
            matriz = self._leer_excel(contenido)
        else:
            raise ArchivoInvalidoException(
                f"Extensión '.{extension}' no soportada. Solo se aceptan .csv y .xlsx."
            )
        return self._matriz_a_filas(matriz)

    def _leer_csv(self, contenido: bytes) -> List[list]:
        try:
            texto = contenido.decode("utf-8-sig")
        except UnicodeDecodeError:
            texto = contenido.decode("latin-1")  # CSV guardados como "ANSI" desde Excel
        primera = texto.split("\n", 1)[0]
        delimitador = ";" if primera.count(";") > primera.count(",") else ","  # Excel en español usa ;
        return list(csv.reader(io.StringIO(texto), delimiter=delimitador))

    def _leer_excel(self, contenido: bytes) -> List[list]:
        try:
            libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        except Exception:
            raise ArchivoInvalidoException("El archivo no es un Excel (.xlsx) válido.")
        try:
            return [list(f) for f in libro.active.iter_rows(values_only=True)]
        finally:
            libro.close()

    def _matriz_a_filas(self, matriz: List[list]) -> List[Tuple[int, Dict[str, Any]]]:
        if not matriz:
            raise ArchivoInvalidoException("El archivo está vacío.")
        encabezados = [str(c).strip().lower() if c is not None else "" for c in matriz[0]]
        faltantes = COLUMNAS_OBLIGATORIAS - set(encabezados)
        if faltantes:
            raise ArchivoInvalidoException(f"Faltan columnas obligatorias: {', '.join(sorted(faltantes))}.")

        filas = []
        for nro, fila in enumerate(matriz[1:], start=2):  # fila 1 = encabezado
            if all(v is None or str(v).strip() == "" for v in fila):
                continue  # ignora filas vacías (conservando la numeración real del archivo)
            filas.append((nro, {enc: (fila[i] if i < len(fila) else None) for i, enc in enumerate(encabezados)}))
        return filas

    # ---------- Validación fila por fila ----------
    def _validar_fila(self, fila: dict, ids_validos: set) -> Tuple[Optional[dict], Optional[str]]:
        """Retorna (dict_limpio, None) si la fila es válida, o (None, 'motivo') si no."""
        try:
            numero = _a_float(fila["id_estudiante"])
            if not numero.is_integer():
                raise ValueError
            id_estudiante = int(numero)
        except (TypeError, ValueError, KeyError):
            return None, "id_estudiante no es un número entero"

        if id_estudiante not in ids_validos:
            return None, f"id_estudiante {id_estudiante} no pertenece a este curso"

        fecha_raw = fila.get("fecha")
        try:
            fecha = fecha_raw.strftime("%Y-%m-%d") if hasattr(fecha_raw, "strftime") else str(fecha_raw).strip()[:10]
            datetime.strptime(fecha, "%Y-%m-%d")
        except (ValueError, TypeError):
            return None, "fecha inválida (formato esperado AAAA-MM-DD)"

        try:
            nota = _a_float(fila["nota"])
        except (TypeError, ValueError):
            return None, "nota no es un número"
        if not (1.0 <= nota <= 7.0):
            return None, f"nota {nota} fuera de rango (debe ser entre 1.0 y 7.0)"

        try:
            asistencia = _a_float(fila["asistencia"])
        except (TypeError, ValueError):
            return None, "asistencia no es un número"
        if not (0.0 <= asistencia <= 100.0):
            return None, f"asistencia {asistencia} fuera de rango (debe ser entre 0 y 100)"

        return {"id_estudiante": id_estudiante, "fecha": fecha, "nota": nota, "asistencia": asistencia}, None

    # ---------- Orquestación: leer + validar + insertar ----------
    def procesar_carga(self, curso_id: int, nombre_archivo: str, contenido: bytes) -> dict:
        filas = self.leer_archivo(nombre_archivo, contenido)  # puede lanzar ArchivoInvalidoException

        with conexion_bd() as con:
            if not con.execute("SELECT 1 FROM Curso WHERE id = ?", (curso_id,)).fetchone():
                raise CursoNoEncontradoException(f"El curso {curso_id} no existe.")
            ids_validos = {r["id"] for r in con.execute(
                "SELECT id FROM Estudiante WHERE id_curso = ?", (curso_id,)).fetchall()}

            errores, validas, vistos = [], [], {}
            for nro, fila in filas:
                limpia, motivo = self._validar_fila(fila, ids_validos)
                if not motivo:
                    clave = (limpia["id_estudiante"], limpia["fecha"])
                    if clave in vistos:
                        motivo = f"duplicada en el archivo (mismo alumno y fecha que la fila {vistos[clave]})"
                    else:
                        vistos[clave] = nro
                if motivo:
                    errores.append({"fila": nro, "motivo": motivo})
                else:
                    validas.append(limpia)

            if validas:
                # Reemplaza lo anterior del mismo alumno+fecha (registros generales, sin asignatura)
                con.executemany(
                    "DELETE FROM RegistroAcademico WHERE id_estudiante = ? AND fecha = ? AND id_asignatura IS NULL",
                    [(f["id_estudiante"], f["fecha"]) for f in validas],
                )
                con.executemany(
                    "INSERT INTO RegistroAcademico (id_estudiante, fecha, asistencia, nota) VALUES (?,?,?,?)",
                    [(f["id_estudiante"], f["fecha"], f["asistencia"], f["nota"]) for f in validas],
                )

        return {
            "exito": True,
            "filas_procesadas": len(filas),
            "filas_insertadas": len(validas),
            "filas_con_error": len(errores),
            "errores": errores,
        }
