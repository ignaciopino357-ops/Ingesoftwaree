"""Módulo Gestión Académica · ADAPTADOR DE SALIDA: planillas Excel con openpyxl (HU-08)."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.academico.dominio import NUM_EVALUACIONES, Alumno
from app.academico.puertos import ExportadorPlanillas

ENCABEZADO = Font(bold=True, color="FFFFFF")
FONDO = PatternFill("solid", fgColor="35577D")


def _formatear(ws) -> None:
    for celda in ws[1]:
        celda.font = ENCABEZADO
        celda.fill = FONDO
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions["A"].width = 28
    ws.freeze_panes = "B2"


def _a_bytes(wb: Workbook) -> bytes:
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


class ExportadorExcel(ExportadorPlanillas):
    def notas(self, alumnos: list[Alumno], hojas: list[tuple[str, dict]]) -> bytes:
        """Una hoja por asignatura: fila por alumno, columna por evaluación y promedio."""
        wb = Workbook()
        wb.remove(wb.active)
        for asignatura, notas in hojas:
            ws = wb.create_sheet(asignatura[:31])
            ws.append(["Alumno"] + [f"N{i}" for i in range(1, NUM_EVALUACIONES + 1)] + ["Promedio"])
            for alumno in alumnos:
                propias = notas.get(alumno.id, {})
                valores = list(propias.values())
                ws.append([alumno.nombre]
                          + [propias.get(i) for i in range(1, NUM_EVALUACIONES + 1)]
                          + [round(sum(valores) / len(valores), 1) if valores else None])
            _formatear(ws)
        return _a_bytes(wb)

    def asistencia(self, alumnos: list[Alumno], fechas: list[str], registros: dict, conteos: dict) -> bytes:
        """Fila por alumno, columna por día (P, A o J) y totales al final."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Asistencia"
        ws.append(["Alumno"] + [f"{f[8:10]}-{f[5:7]}" for f in fechas]
                  + ["Presente", "Ausente", "Justificado", "% Asistencia"])
        for alumno in alumnos:
            conteo = conteos[alumno.id]
            total = sum(conteo.values())
            porcentaje = round((conteo["P"] + conteo["J"]) * 100 / total, 1) if total else None
            ws.append([alumno.nombre] + [registros[alumno.id].get(f, "") for f in fechas]
                      + [conteo["P"], conteo["A"], conteo["J"], porcentaje])
        _formatear(ws)
        for i in range(2, len(fechas) + 2):
            ws.column_dimensions[get_column_letter(i)].width = 6
        ws.append([])
        ws.append(["P = Presente · A = Ausente · J = Justificado"])
        return _a_bytes(wb)
