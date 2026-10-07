"""Pruebas de los criterios de aceptación de los Sprints 1 y 2 a través de la web."""
from io import BytesIO

from openpyxl import load_workbook

from tests.conftest import entrar


def _id_curso(cliente, nombre="3° Medio A"):
    # pjara tiene un solo curso: /profesor redirige a /profesor/cursos/<id>
    r = cliente.get("/profesor", follow_redirects=False)
    return int(r.headers["location"].rsplit("/", 1)[1])


# ---------- HU-01 ----------

def test_cada_rol_llega_a_su_portal(cliente):
    for usuario, clave, portal in [("directiva", "directiva123", "/directiva"),
                                   ("pjara", "profe123", "/profesor"),
                                   ("alumno1", "alumno123", "/alumno")]:
        r = entrar(cliente, usuario, clave)
        assert r.status_code == 303 and r.headers["location"] == portal
        cliente.post("/logout")


def test_clave_incorrecta(cliente):
    r = entrar(cliente, "pjara", "mala")
    assert r.status_code == 401
    assert "Usuario o contraseña incorrectos" in r.text


def test_sin_sesion_va_al_login(cliente):
    r = cliente.get("/profesor", follow_redirects=False)
    assert r.headers["location"] == "/login"


def test_otro_portal_bloqueado(cliente):
    entrar(cliente, "alumno1", "alumno123")
    r = cliente.get("/directiva")
    assert r.status_code == 403
    assert "No tiene acceso a esta sección" in r.text


# ---------- HU-04 + HU-07 ----------

def test_profesor_ve_su_curso_ordenado_por_riesgo(cliente):
    entrar(cliente, "pjara", "profe123")
    r = cliente.get("/profesor", follow_redirects=True)
    assert "3° Medio A" in r.text
    # el alumno de riesgo Alto aparece antes que los de riesgo Bajo
    assert r.text.index("Emilia Vargas") < r.text.index("Ana Contreras")
    assert "Sin datos" in r.text


def test_profesor_no_ve_curso_ajeno(cliente):
    entrar(cliente, "cmunoz", "profe123")
    propio = _id_curso(cliente)
    r = cliente.get(f"/profesor/cursos/{propio - 1}")
    assert r.status_code == 403
    assert "No tiene acceso a este curso" in r.text


# ---------- HU-05 ----------

def test_nota_invalida_no_guarda_nada(cliente):
    entrar(cliente, "pjara", "profe123")
    curso = _id_curso(cliente)
    pagina = cliente.get(f"/profesor/cursos/{curso}/notas?asignatura=1")
    assert pagina.status_code == 200
    r = cliente.post(f"/profesor/cursos/{curso}/notas",
                     data={"asignatura_id": "1", "n_8_1": "6,5", "n_8_2": "8.0"})
    assert r.status_code == 400
    assert "Debe estar entre 1.0 y 7.0" in r.text
    # la nota 6,5 tampoco quedó guardada
    alumno = cliente.get(f"/profesor/cursos/{curso}/alumnos/8")
    assert "6.5" not in alumno.text


def test_guardar_notas_actualiza_riesgo(cliente):
    entrar(cliente, "pjara", "profe123")
    curso = _id_curso(cliente)
    # Hugo Navarro (id 12) no tiene datos: con notas bajas pasa a Medio por promedio
    datos = {"asignatura_id": "1", "n_12_1": "3,0", "n_12_2": "3.2"}
    r = cliente.post(f"/profesor/cursos/{curso}/notas", data=datos, follow_redirects=True)
    assert "Notas guardadas" in r.text
    ficha = cliente.get(f"/profesor/cursos/{curso}/alumnos/12")
    assert "Medio" in ficha.text and "Promedio bajo 4.0" in ficha.text


# ---------- HU-06 ----------

def test_asistencia_bloquea_feriado_y_fin_de_semana(cliente):
    entrar(cliente, "pjara", "profe123")
    curso = _id_curso(cliente)
    r = cliente.get(f"/profesor/cursos/{curso}/asistencia?fecha=2026-09-18")
    assert "Independencia Nacional" in r.text
    r = cliente.get(f"/profesor/cursos/{curso}/asistencia?fecha=2026-10-03")
    assert "fin de semana" in r.text
    r = cliente.post(f"/profesor/cursos/{curso}/asistencia",
                     data={"fecha": "2026-10-03", "estado_12": "A"}, follow_redirects=True)
    assert "No se guardó" in r.text


def test_pasar_lista_guarda_estados(cliente):
    entrar(cliente, "pjara", "profe123")
    curso = _id_curso(cliente)
    r = cliente.post(f"/profesor/cursos/{curso}/asistencia",
                     data={"fecha": "2026-10-07", "estado_12": "A", "estado_6": "J"},
                     follow_redirects=True)
    assert "guardada" in r.text
    ficha = cliente.get(f"/profesor/cursos/{curso}/alumnos/12")
    assert "Ausente 1" in ficha.text


# ---------- HU-08 ----------

def test_exportar_asistencia_por_dia(cliente):
    entrar(cliente, "pjara", "profe123")
    curso = _id_curso(cliente)
    r = cliente.get(f"/profesor/cursos/{curso}/exportar/asistencia")
    assert r.status_code == 200
    ws = load_workbook(BytesIO(r.content)).active
    encabezado = [c.value for c in ws[1]]
    assert encabezado[0] == "Alumno"
    assert encabezado[-4:] == ["Presente", "Ausente", "Justificado", "% Asistencia"]
    valores = {c.value for fila in ws.iter_rows(min_row=2, max_row=9) for c in fila[1:-4]}
    assert {"P", "A", "J"} <= valores


def test_exportar_notas_una_hoja_por_asignatura(cliente):
    entrar(cliente, "pjara", "profe123")
    curso = _id_curso(cliente)
    r = cliente.get(f"/profesor/cursos/{curso}/exportar/notas")
    wb = load_workbook(BytesIO(r.content))
    assert "Lenguaje" in wb.sheetnames and len(wb.sheetnames) == 5


def test_exportar_sin_datos(cliente):
    entrar(cliente, "directiva", "directiva123")
    cliente.post("/directiva/cursos", data={"nombre": "4° Medio C", "profesor_id": "4"})
    cliente.post("/logout")
    entrar(cliente, "pnuevo", "profe123")
    curso = _id_curso(cliente)
    r = cliente.get(f"/profesor/cursos/{curso}/exportar/notas", follow_redirects=True)
    assert "No hay datos para exportar" in r.text


# ---------- HU-02 ----------

def test_directiva_asigna_alumno_sin_curso(cliente):
    entrar(cliente, "directiva", "directiva123")
    r = cliente.post("/directiva/cursos/1/alumnos", data={"alumno_id": "15"}, follow_redirects=True)
    assert "Alumno agregado" in r.text


def test_alumno_no_puede_estar_en_dos_cursos(cliente):
    entrar(cliente, "directiva", "directiva123")
    # alumno1 (id 5) está en 3° Medio A; se intenta agregar a 3° Medio B (id 2)
    r = cliente.post("/directiva/cursos/2/alumnos", data={"alumno_id": "5"}, follow_redirects=True)
    assert "ya pertenece a 3° Medio A" in r.text


# ---------- HU-03 ----------

def test_semestre_con_fechas_invertidas(cliente):
    entrar(cliente, "directiva", "directiva123")
    r = cliente.post("/directiva/calendario/semestre",
                     data={"inicio": "2026-12-01", "fin": "2026-08-01"}, follow_redirects=True)
    assert "término es anterior" in r.text


def test_agregar_feriado_bloquea_asistencia(cliente):
    entrar(cliente, "directiva", "directiva123")
    cliente.post("/directiva/calendario/feriados", data={"fecha": "2026-10-06", "nombre": "Feriado de prueba"})
    cal = cliente.get("/directiva/calendario?mes=2026-10")
    assert "Feriado de prueba" in cal.text
    cliente.post("/logout")
    entrar(cliente, "pjara", "profe123")
    curso = _id_curso(cliente)
    r = cliente.get(f"/profesor/cursos/{curso}/asistencia?fecha=2026-10-06")
    assert "Feriado de prueba" in r.text
