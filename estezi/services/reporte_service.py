# API/services/reporte_service.py
# Tarea (Back): servicio de generación de archivos (reporte tabular por curso)
# Usa openpyxl para .xlsx y csv para .csv (sin dependencias externas pesadas).

import csv
import io
import re
import unicodedata
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from database import obtener_conexion


class ReporteService:

    def _obtener_datos_curso(self, curso_id: int):
        con = obtener_conexion()
        curso = con.execute("SELECT * FROM Curso WHERE id = ?", (curso_id,)).fetchone()
        if not curso:
            con.close()
            return None, None
        filas = con.execute(
            """
            SELECT e.nombres, e.apellidos, e.rut, e.nivel_riesgo_actual,
                   ROUND(AVG(ra.nota), 1)       AS promedio_notas,
                   ROUND(AVG(ra.asistencia), 1) AS promedio_asistencia
            FROM Estudiante e
            LEFT JOIN RegistroAcademico ra ON ra.id_estudiante = e.id
            WHERE e.id_curso = ?
            GROUP BY e.id
            ORDER BY e.apellidos
            """,
            (curso_id,),
        ).fetchall()
        con.close()
        return dict(curso), [dict(f) for f in filas]

    ENCABEZADOS = ["Nombres", "Apellidos", "RUT", "Promedio Notas", "Promedio Asistencia (%)", "Nivel de Riesgo"]

    def _nombre_archivo(self, nombre_curso: str, extension: str) -> str:
        """Nombre de archivo seguro para el header HTTP (sin tildes, ° ni espacios)."""
        sin_tildes = unicodedata.normalize("NFKD", nombre_curso).encode("ascii", "ignore").decode()
        limpio = re.sub(r"[^A-Za-z0-9]+", "_", sin_tildes).strip("_")
        return f"reporte_{limpio}.{extension}"

    def _fila(self, e: dict):
        return [e["nombres"], e["apellidos"], e["rut"], e["promedio_notas"], e["promedio_asistencia"],
                e["nivel_riesgo_actual"]]

    def generar_csv(self, curso_id: int):
        curso, estudiantes = self._obtener_datos_curso(curso_id)
        if curso is None:
            return None, None
        buffer = io.StringIO()
        escritor = csv.writer(buffer)
        escritor.writerow([f"Reporte de curso: {curso['nombre']}"])
        escritor.writerow([f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}"])
        escritor.writerow([])
        escritor.writerow(self.ENCABEZADOS)
        for e in estudiantes:
            escritor.writerow(self._fila(e))
        nombre = self._nombre_archivo(curso["nombre"], "csv")
        return buffer.getvalue().encode("utf-8-sig"), nombre

    def generar_excel(self, curso_id: int):
        curso, estudiantes = self._obtener_datos_curso(curso_id)
        if curso is None:
            return None, None

        libro = Workbook()
        hoja = libro.active
        hoja.title = "Reporte de curso"

        hoja.merge_cells("A1:F1")
        hoja["A1"] = f"Reporte de curso: {curso['nombre']}"
        hoja["A1"].font = Font(bold=True, size=13)
        hoja["A2"] = f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}"

        fila_encabezado = 4
        relleno = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
        for col, titulo in enumerate(self.ENCABEZADOS, start=1):
            celda = hoja.cell(row=fila_encabezado, column=col, value=titulo)
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = relleno

        for i, e in enumerate(estudiantes, start=fila_encabezado + 1):
            for col, valor in enumerate(self._fila(e), start=1):
                hoja.cell(row=i, column=col, value=valor)

        for col, ancho in zip("ABCDEF", [16, 16, 14, 14, 20, 14]):
            hoja.column_dimensions[col].width = ancho

        buffer = io.BytesIO()
        libro.save(buffer)
        nombre = self._nombre_archivo(curso["nombre"], "xlsx")
        return buffer.getvalue(), nombre
