"""HU-08: exportación a Excel de notas y asistencia día a día."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app import consultas

ENCABEZADO = Font(bold=True, color="FFFFFF")
FONDO = PatternFill("solid", fgColor="35577D")


def _formatear(ws, anchos: dict) -> None:
    for celda in ws[1]:
        celda.font = ENCABEZADO
        celda.fill = FONDO
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, ancho in anchos.items():
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "B2"


def _a_bytes(wb: Workbook) -> bytes:
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


def excel_notas(curso_id: int) -> bytes:
    """Una hoja por asignatura: fila por alumno, columna por evaluación y promedio."""
    wb = Workbook()
    wb.remove(wb.active)
    alumnos = consultas.alumnos_del_curso(curso_id)
    for asig in consultas.asignaturas():
        ws = wb.create_sheet(asig["nombre"][:31])
        ws.append(["Alumno"] + [f"N{i}" for i in range(1, consultas.NUM_EVALUACIONES + 1)] + ["Promedio"])
        notas = consultas.notas_de_asignatura(curso_id, asig["id"])
        for alumno in alumnos:
            propias = notas.get(alumno["id"], {})
            fila = [alumno["nombre"]]
            fila += [propias.get(i) for i in range(1, consultas.NUM_EVALUACIONES + 1)]
            valores = list(propias.values())
            fila.append(round(sum(valores) / len(valores), 1) if valores else None)
            ws.append(fila)
        _formatear(ws, {"A": 28})
    return _a_bytes(wb)


def excel_asistencia(curso_id: int) -> bytes:
    """Fila por alumno, columna por día (P, A o J) y totales al final."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Asistencia"
    fechas = consultas.fechas_con_asistencia(curso_id)
    encabezado = ["Alumno"] + [f"{f[8:10]}-{f[5:7]}" for f in fechas]
    encabezado += ["Presente", "Ausente", "Justificado", "% Asistencia"]
    ws.append(encabezado)

    for alumno in consultas.alumnos_del_curso(curso_id):
        por_fecha = {r["fecha"]: r["estado"] for r in consultas.asistencia_del_alumno(alumno["id"])}
        conteo = consultas.conteo_asistencia(alumno["id"])
        total = sum(conteo.values())
        porcentaje = round((conteo["P"] + conteo["J"]) * 100 / total, 1) if total else None
        ws.append(
            [alumno["nombre"]]
            + [por_fecha.get(f, "") for f in fechas]
            + [conteo["P"], conteo["A"], conteo["J"], porcentaje]
        )

    _formatear(ws, {"A": 28})
    for i in range(2, len(fechas) + 2):
        ws.column_dimensions[get_column_letter(i)].width = 6
    ws.append([])
    ws.append(["P = Presente · A = Ausente · J = Justificado"])
    return _a_bytes(wb)
